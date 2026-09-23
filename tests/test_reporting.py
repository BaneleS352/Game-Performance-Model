from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from game_performance.reporting import (
    BUSINESS_REPORT_EXCLUDED_COLUMNS,
    build_calibration_report,
    build_forecast_capability_report,
    build_holdout_comparison,
    export_business_report,
    export_excel,
)


def test_forecast_capability_reports_history_horizon_games_and_average_error():
    comparison = build_holdout_comparison(
        pd.DataFrame(
            {
                "Game_ID": ["A", "B"],
                "First_Observed_Date": ["2026-01-01", "2026-01-02"],
                "Obs_Window": [21, 21],
                "Horizon": [7, 7],
                "actual_clients": [100, 200],
                "predicted_clients": [90, 210],
                "predicted_clients_low": [80, 180],
                "predicted_clients_high": [110, 220],
                "actual_stakes": [1_000, 2_000],
                "predicted_stakes": [900, 2_100],
                "predicted_stakes_low": [800, 1_800],
                "predicted_stakes_high": [1_100, 2_200],
                "actual_betcount": [50, 100],
                "predicted_betcount": [55, 90],
                "predicted_betcount_low": [45, 80],
                "predicted_betcount_high": [60, 110],
                "actual_ggr": [10, 20],
                "predicted_ggr": [8, 25],
                "predicted_ggr_low": [5, 12],
                "predicted_ggr_high": [12, 18],
            }
        )
    )

    report = build_forecast_capability_report(comparison)

    assert report.loc[0, "History_Days"] == 21
    assert report.loc[0, "Forecast_Days"] == 7
    assert report.loc[0, "Games_Evaluated"] == 2
    assert report.loc[0, "Average_Off_Percent"] == pytest.approx(11.6667, rel=1e-4)
    assert report.loc[0, "Prediction_Range_Hit_Percent"] == 87.5
    assert "11.7% average error" in report.loc[0, "Summary"]


def test_holdout_comparison_uses_next_seven_day_probability_and_off_flags():
    holdout = pd.DataFrame(
        {
            "Game_ID": [1],
            "First_Observed_Date": ["2026-01-01"],
            "Obs_Window": [14],
            "Horizon": [14],
            "actual_clients": [100.0],
            "predicted_clients_low": [70.0],
            "predicted_clients": [85.0],
            "predicted_clients_high": [105.0],
            "confidence_clients": [0.75],
        }
    )

    comparison = build_holdout_comparison(holdout)

    expected = [
        "actual_clients",
        "predicted_clients_low",
        "predicted_clients",
        "predicted_clients_high",
        "confidence_clients",
        "variance_clients",
        "probability_clients",
        "clients_actual_in_range",
        "clients_range_width",
    ]
    start = comparison.columns.get_loc("actual_clients")

    assert list(comparison.columns[start : start + len(expected)]) == expected
    assert comparison.loc[0, "variance_clients"] == -15.0
    assert comparison.loc[0, "probability_clients"] == 0.85
    assert comparison.loc[0, "clients_actual_in_range"] == 1
    assert comparison.loc[0, "clients_range_width"] == 35.0
    assert comparison.loc[0, "off_by_10_pct"] == 1
    assert comparison.loc[0, "off_by_20_pct"] == 0
    assert "accuracy_clients" not in comparison.columns
    assert "absolute_error_clients" not in comparison.columns
    assert "variance_pct_clients" not in comparison.columns
    assert "High_Performer_Match" not in comparison.columns


def test_holdout_comparison_keeps_missing_range_actuals_missing():
    holdout = pd.DataFrame(
        {
            "Game_ID": [1],
            "First_Observed_Date": ["2026-01-01"],
            "Obs_Window": [14],
            "Horizon": [14],
            "actual_clients": [None],
            "predicted_clients_low": [70.0],
            "predicted_clients": [85.0],
            "predicted_clients_high": [105.0],
        }
    )

    comparison = build_holdout_comparison(holdout)

    assert pd.isna(comparison.loc[0, "clients_actual_in_range"])


def test_calibration_report_buckets_confidence_against_actual_rate():
    comparison = pd.DataFrame(
        {
            "is_high_performer": [0, 1, 1],
            "probability_high_performance": [0.05, 0.72, 0.76],
            "raw_probability_high_performance": [0.15, 0.90, 0.95],
        }
    )

    report = build_calibration_report(comparison, bins=2)

    assert list(report["Probability_Bin"]) == ["0.0-0.5", "0.5-1.0"]
    assert list(report["Count"]) == [1, 2]
    assert report.loc[1, "Actual_High_Performer_Rate"] == 1.0
    assert "Mean_Raw_Probability" in report.columns


def test_export_excel_keeps_all_regression_metrics_and_calibration_sheets():
    comparison = pd.DataFrame(
        {
            "Game_ID": [1, 2],
            "predicted_high_performer": [0, 1],
            "actual_clients": [10, 20],
            "predicted_clients": [11, 19],
            "actual_stakes": [100, 200],
            "predicted_stakes": [110, 190],
        }
    )

    path = Path("artifacts/test_report.xlsx")
    workbook = None
    try:
        export_excel(
            path,
            comparison=comparison,
            forecasts=pd.DataFrame(
                {
                    "Game_ID": [1],
                    "recommendation": ["Promote"],
                    "Best_Iteration": [42],
                    "predicted_clients_low": [8],
                    "predicted_clients": [10],
                    "predicted_clients_high": [12],
                    "confidence_clients": [0.8],
                    "predicted_stakes_low": [90],
                    "predicted_stakes": [100],
                    "predicted_stakes_high": [110],
                    "confidence_stakes": [0.7],
                    "probability_high_client_performer": [0.8],
                    "predicted_high_client_performer": [1],
                    "probability_high_stakes_performer": [0.7],
                    "predicted_high_stakes_performer": [1],
                    "predicted_ggr": [5],
                }
            ),
            regression=pd.DataFrame(
                {
                    "Target": ["target_clients", "target_ggr", "target_stakes"],
                    "MAE": [1.0, 2.0, 3.0],
                }
            ),
            classification=pd.DataFrame(),
            uncertainty_validation=pd.DataFrame(
                {
                    "Target": ["target_stakes"],
                    "Segment_Type": ["Provider"],
                    "Observed_Coverage": [0.8],
                }
            ),
        )

        workbook = load_workbook(path, read_only=True)

        assert workbook.sheetnames == [
            "Forecast Capability",
            "Holdout Comparison",
            "High Performers",
            "Forecasts",
            "Model Accuracy",
            "Promo Impact Periods",
            "Uncertainty Validation",
            "future_performance_score_calcul",
        ]

        forecast_headers = next(workbook["Forecasts"].iter_rows(values_only=True))
        assert forecast_headers == (
            "Game_ID",
            "predicted_clients_low",
            "predicted_clients",
            "predicted_clients_high",
            "confidence_clients",
            "predicted_ggr",
            "predicted_stakes_low",
            "predicted_stakes",
            "predicted_stakes_high",
            "confidence_stakes",
            "probability_high_client_performer",
            "predicted_high_client_performer",
            "probability_high_stakes_performer",
            "predicted_high_stakes_performer",
        )

        accuracy_rows = list(workbook["Model Accuracy"].iter_rows(values_only=True))
        assert accuracy_rows[0] == ("Target", "MAE")
        assert accuracy_rows[1] == ("target_clients", 1.0)
        assert accuracy_rows[2] == ("target_ggr", 2.0)
        assert accuracy_rows[3] == ("target_stakes", 3.0)

        score_rows = list(workbook["future_performance_score_calcul"].iter_rows(values_only=True))
        assert score_rows[0][:3] == ("Metric", "Target field", "Weight")
        assert score_rows[1][:2] == ("High client performer", "is_high_client_performer")
        assert score_rows[2][:2] == ("High stakes performer", "is_high_stakes_performer")
        assert score_rows[3][:2] == ("Balanced high performer", "is_balanced_high_performer")
    finally:
        if workbook is not None:
            workbook.close()
        path.unlink(missing_ok=True)


def test_business_report_has_only_concise_sheets_and_excludes_technical_columns():
    path = Path("artifacts/test_business_report.xlsx")
    workbook = None
    comparison = pd.DataFrame(
        {
            "Game_ID": ["holdout-game"],
            "TournamentName": ["Game"],
            "SystemName": ["System"],
            "First_Observed_Date": ["2026-01-01"],
            "Observation_Start": ["2026-01-01"],
            "Observation_End": ["2026-01-14"],
            "Prediction_Start": ["2026-01-15"],
            "Prediction_End": ["2026-01-21"],
            "future_performance_score": [0.9],
            "is_high_client_performer": [1],
            "is_high_stakes_performer": [1],
            "is_balanced_high_performer": [1],
            "predicted_high_performer": [1],
        }
    )
    forecasts = pd.DataFrame(
        {
            "Game_ID": ["forecast-game"],
            "TournamentName": ["Game"],
            "SystemName": ["System"],
            "First_Observed_Date": ["2026-01-01"],
            "predicted_clients": [10],
            "predicted_stakes": [100],
            "raw_probability_high_performance": [0.9],
            "probability_high_performance": [0.9],
            "classification_threshold": [0.8],
            "predicted_high_performer": [1],
        }
    )

    try:
        export_business_report(path, comparison=comparison, forecasts=forecasts)
        workbook = load_workbook(path, read_only=True)

        assert workbook.sheetnames == ["Forecast Capability", "Holdout Comparison", "High Performers", "Forecasting", "Promo Impact Periods"]
        for sheet_name in workbook.sheetnames:
            headers = set(next(workbook[sheet_name].iter_rows(values_only=True)))
            if sheet_name != "Promo Impact Periods":
                assert not headers.intersection(BUSINESS_REPORT_EXCLUDED_COLUMNS)
        assert workbook["High Performers"].max_row == 2
    finally:
        if workbook is not None:
            workbook.close()
            path.unlink(missing_ok=True)


def test_exports_external_impact_flags_and_periods():
    path = Path("artifacts/test_impact_report.xlsx")
    workbook = None
    periods = pd.DataFrame(
        {
            "Game_ID": ["forecast-game"],
            "SystemName": ["System"],
            "TournamentName": ["Game"],
            "Start_Date": ["2026-01-04"],
            "End_Date": ["2026-01-04"],
            "Metric": ["Clients"],
            "Direction": ["Spike"],
            "Baseline_Value": [10.0],
            "Observed_Value": [20.0],
            "Percent_Change": [100.0],
            "Severity": ["High"],
            "Reason": ["Inferred external impact"],
        }
    )
    forecasts = pd.DataFrame(
        {
            "Game_ID": ["forecast-game"],
            "external_impact_detected": [True],
            "external_impact_periods": ["2026-01-04 to 2026-01-04 (Clients spike)"],
            "external_impact_reason": ["Inferred external impact"],
        }
    )

    try:
        export_excel(
            path,
            comparison=pd.DataFrame(),
            forecasts=forecasts,
            regression=pd.DataFrame(),
            classification=pd.DataFrame(),
            impact_periods=periods,
        )
        workbook = load_workbook(path, read_only=True)
        headers = next(workbook["Forecasts"].iter_rows(values_only=True))
        assert "external_impact_detected" in headers
        assert list(workbook["Promo Impact Periods"].iter_rows(values_only=True))[1][0] == "forecast-game"
    finally:
        if workbook is not None:
            workbook.close()
        path.unlink(missing_ok=True)
